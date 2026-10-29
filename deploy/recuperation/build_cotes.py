"""Liste et étiquettes QR d'une plage de cartons cotés (ex. 2010-KOU-0001 à 0320).
usage: python3 build_cotes.py cartons.json sortie_dossier"""
import json, sys, collections, datetime, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.graphics.barcode import qr
from reportlab.graphics.shapes import Drawing
from reportlab.graphics import renderPDF

src, outdir = sys.argv[1], sys.argv[2]
data = json.load(open(src, encoding="utf-8"))
APP = "https://geoarchives.ceiba-analytics.com"
cartons = data["cartons"]


def tidy(text):
    return " ".join(str(text or "").split())


def pdf_text(text):
    """Les polices PDF de base ne couvrent que le latin-1 : apostrophes et tirets typographiques remplacés."""
    text = tidy(text).replace("’", "'").replace("‘", "'").replace("–", "-").replace("—", "-").replace(" ", " ")
    return "".join(ch if ord(ch) < 256 else unicodedata.normalize("NFKD", ch).encode("latin-1", "ignore").decode("latin-1") for ch in text)


for c in cartons:
    refs = collections.Counter(tidy(d["referenceClassement"]) for d in c["dossiers"] if tidy(d["referenceClassement"]))
    c["reference"] = refs.most_common(1)[0][0] if refs else None
    c["nom"] = c["listeLibelle"] or c["reference"] or (collections.Counter(tidy(d["libelle"]) for d in c["dossiers"] if tidy(d["libelle"])).most_common(1) or [("Sans libellé", 0)])[0][0]
    first = c["dossiers"][0] if c["dossiers"] else {}
    c["lieu"] = " · ".join(x for x in [first.get("site"), first.get("salle")] if x)

total_dossiers = sum(len(c["dossiers"]) for c in cartons)
liste = [c for c in cartons if c["listeRef"] is not None]
vides = [c for c in cartons if not c["dossiers"]]
tag = f'{data["prefixe"]}-{data["de"]:04d}-{data["a"]:04d}'

# ---------- Excel ----------
wb = Workbook()
H = Font(name="Arial", bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="17284D")
F = Font(name="Arial", size=10); WARN = PatternFill("solid", fgColor="FDECC8")
thin = Side(style="thin", color="D0D5DD")


def sheet(ws, headers, widths):
    ws.append(headers)
    for cell, width in zip(ws[1], widths):
        cell.font = H; cell.fill = HF; cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 32


def finish(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = F; cell.border = Border(bottom=thin); cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions


ws = wb.active; ws.title = "Cartons"
sheet(ws, ["N°", "Cote du carton", "Réf. provisoire (liste)", "Libellé de la liste", "Référence de classement (base)", "Dossiers", "Cotes des dossiers", "Site · salle", "Remarque"],
      [6, 17, 12, 26, 30, 10, 30, 32, 44])
for i, c in enumerate(cartons, 1):
    n = len(c["dossiers"])
    remarque = "" if n else "Aucun dossier en base : rattacher les dossiers à ce carton lors de l'inventaire (scan du QR)."
    ws.append([i, c["cote"], c["listeRef"], c["listeLibelle"], c["reference"], n,
               f'{c["cote"]}-001 à {c["cote"]}-{n:03d}' if n else "", c["lieu"], remarque])
    if not n:
        for cell in ws[ws.max_row]: cell.fill = WARN
finish(ws)

wd = wb.create_sheet("Dossiers")
sheet(wd, ["Cote du dossier", "Carton", "N° dans le carton", "Titulaire", "Îlot", "Lot", "Titre foncier", "Lotissement", "Commune", "Référence de classement", "Référence dossier"],
      [22, 17, 10, 34, 8, 10, 14, 30, 14, 26, 18])
for c in cartons:
    for d in c["dossiers"]:
        wd.append([d["cote"], c["cote"], d["rang"], d["titulaire"], d["ilot"], d["lot"], d["titreFoncier"], d["lotissement"], d["commune"], d["referenceClassement"], d["reference"]])
finish(wd)

wn = wb.create_sheet("Lisez-moi")
for line in [
    ["Liste et étiquettes QR des cartons", ""],
    ["Cartons", f"{len(cartons)} : {cartons[0]['cote']} à {cartons[-1]['cote']}" if cartons else "0"],
    ["Dossiers", total_dossiers],
    ["Cote d'un carton", "2010-KOU-0001 : rubrique 2010 (Dossiers ACD), commune (KOU = Koumassi), n° du carton."],
    ["Cote d'un dossier", "Cote du carton suivie du n° du dossier : 2010-KOU-0001-001, -002… sans trou."],
    ["QR codes", "Le QR de l'étiquette ouvre la fiche du carton dans GeoArchives (connexion requise)."],
    ["Liste d'enlèvement", f"Les {len(liste)} premiers cartons suivent l'ordre des références provisoires 1 à {len(liste)} de la liste d'enlèvement."],
    ["Données au", data["exportedAt"]],
]:
    wn.append(line)
wn.column_dimensions["A"].width = 24; wn.column_dimensions["B"].width = 100
for cell in wn["A"]: cell.font = Font(name="Arial", bold=True)
xlsx = f"{outdir}/liste-cartons-{tag}.xlsx"
wb.save(xlsx)


def draw_qr(c, url, x, y, size):
    widget = qr.QrCodeWidget(url, barLevel="M"); b = widget.getBounds()
    drawing = Drawing(size, size, transform=[size / (b[2] - b[0]), 0, 0, size / (b[3] - b[1]), 0, 0]); drawing.add(widget)
    renderPDF.draw(drawing, c, x, y)


def wrap(text, width, font, size, c, max_lines):
    words, lines, line = pdf_text(text).split(), [], ""
    for word in words:
        trial = f"{line} {word}".strip()
        if c.stringWidth(trial, font, size) <= width: line = trial
        else:
            if line: lines.append(line)
            line = word
    if line: lines.append(line)
    if len(lines) > max_lines:
        lines = lines[:max_lines]; lines[-1] = lines[-1][:-3] + "..."
    return lines


# ---------- Étiquettes des cartons (A4, 8 par page, 95 x 66 mm) ----------
def carton_labels(path):
    c = canvas.Canvas(path, pagesize=A4); _, height = A4
    c.setTitle(f"Étiquettes des cartons {tag}")
    cols, rows, lw, lh, mx, my = 2, 4, 95 * mm, 66 * mm, 8 * mm, 12 * mm
    for i, carton in enumerate(cartons):
        if i and i % (cols * rows) == 0: c.showPage()
        k = i % (cols * rows); x = mx + (k % cols) * (lw + 4 * mm); y = height - my - (k // cols + 1) * (lh + 3 * mm)
        c.setStrokeColorRGB(.55, .55, .55); c.roundRect(x, y, lw, lh, 2 * mm)
        draw_qr(c, f"{APP}/c/{carton['cote']}", x + 3 * mm, y + (lh - 42 * mm) / 2, 42 * mm)
        tx, tw = x + 47 * mm, lw - 50 * mm
        c.setFont("Helvetica-Bold", 15); c.drawString(tx, y + lh - 11 * mm, carton["cote"])
        for j, line in enumerate(wrap(carton["nom"], tw, "Helvetica-Bold", 11, c, 3)):
            c.setFont("Helvetica-Bold", 11); c.drawString(tx, y + lh - 19 * mm - j * 5 * mm, line)
        c.setFont("Helvetica", 8.5)
        n = len(carton["dossiers"])
        c.drawString(tx, y + 17 * mm, f"{n} dossier{'s' if n > 1 else ''}" if n else "Dossiers à rattacher")
        if carton["listeRef"] is not None: c.drawString(tx, y + 12.5 * mm, f"Liste d'enlèvement : n° {carton['listeRef']}")
        c.setFont("Helvetica", 7)
        for j, line in enumerate(wrap(carton["lieu"] or "Koumassi", tw, "Helvetica", 7, c, 2)):
            c.drawString(tx, y + 7 * mm - j * 3.2 * mm, line)
    c.save()


# ---------- Étiquettes des dossiers (A4, 3 x 8 par page, 63 x 33 mm) ----------
def dossier_labels(path):
    c = canvas.Canvas(path, pagesize=A4); _, height = A4
    c.setTitle(f"Étiquettes des dossiers {tag}")
    cols, rows, lw, lh, mx, my = 3, 8, 63 * mm, 33 * mm, 7 * mm, 10 * mm
    i = 0
    for carton in cartons:
        n = len(carton["dossiers"])
        for d in carton["dossiers"]:
            if i and i % (cols * rows) == 0: c.showPage()
            k = i % (cols * rows); x = mx + (k % cols) * (lw + 2 * mm); y = height - my - (k // cols + 1) * (lh + 1.5 * mm)
            c.setStrokeColorRGB(.7, .7, .7); c.rect(x, y, lw, lh)
            draw_qr(c, f"{APP}/c/{d['cote']}", x + 1.5 * mm, y + 2.5 * mm, 28 * mm)
            tx, tw = x + 30.5 * mm, lw - 32 * mm
            c.setFont("Helvetica-Bold", 7.6); c.drawString(tx, y + lh - 6 * mm, pdf_text(d["cote"]))
            c.setFont("Helvetica", 7); c.drawString(tx, y + lh - 10 * mm, f"Dossier {d['rang']}/{n}")
            for j, line in enumerate(wrap(d["titulaire"] or "Titulaire non renseigné", tw, "Helvetica-Bold", 7, c, 3)):
                c.setFont("Helvetica-Bold", 7); c.drawString(tx, y + lh - 15 * mm - j * 3.3 * mm, line)
            place = " ".join(x for x in [f"Îlot {d['ilot']}" if d["ilot"] else "", f"lot {d['lot']}" if d["lot"] else ""] if x)
            c.setFont("Helvetica", 6.3)
            if place: c.drawString(tx, y + 3.2 * mm, pdf_text(place)[:30])
            i += 1
    if i == 0:
        c.setFont("Helvetica", 11); c.drawString(20 * mm, height - 30 * mm, "Aucun dossier dans cette plage de cartons.")
    c.save()


carton_labels(f"{outdir}/etiquettes-cartons-{tag}.pdf")
dossier_labels(f"{outdir}/etiquettes-dossiers-{tag}.pdf")
summary = {
    "kind": "cotes", "boxes": len(cartons), "dossiers": total_dossiers,
    "first": cartons[0]["cote"] if cartons else None, "last": cartons[-1]["cote"] if cartons else None,
    "listBoxes": len(liste), "listFound": sum(1 for c in liste if c["dossiers"]), "empty": len(vides),
    "exportedAt": data["exportedAt"],
}
json.dump(summary, open(f"{outdir}/summary.json", "w", encoding="utf-8"), ensure_ascii=False)
print(json.dumps({k: summary[k] for k in ("boxes", "dossiers", "first", "last", "listBoxes", "listFound", "empty")}, ensure_ascii=False))
